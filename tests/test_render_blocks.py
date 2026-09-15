"""V2 阶段一 A1 单测：text 节正文内嵌 markdown 表格的分块渲染（html + docx）。

pytest 运行；也可 python tests/test_render_blocks.py 直接执行。
"""

import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from render.common import split_md_blocks  # noqa: E402

MD_TABLE = """| 指标 | 数值 |
|---|---|
| 储量 | 100万t |


正文段落继续。"""


def _spec_one_text():
    from template_factory.schema import Section, SpecV2
    return SpecV2(report_type="mini", description="d",
                  sections=[Section(id="body", title="正文", kind="text")])


def _doc_one_text(body_text):
    return {"meta": {}, "collections": {}, "facts": []}


def _texts(body_text):
    return [{"section_id": "body", "body": body_text, "cited_fact_ids": []}]


def test_split_blocks_table_and_text():
    blocks = split_md_blocks(f"第一段。\n\n{MD_TABLE}")
    kinds = [k for k, _ in blocks]
    assert kinds == ["text", "table", "text"], kinds
    assert "指标" in blocks[1][1]
    # 单行竖线不成表 → 回退文本
    blocks = split_md_blocks("a | b 单行\n后续")
    assert all(k == "text" for k, _ in blocks)


def test_html_renders_real_table_no_bare_pipes():
    from render.html_report import render
    body = f"段一开头。\n{MD_TABLE}\n段二收尾。"
    html = render(_doc_one_text(body), {"title": "T"}, [], {}, {}, _spec_one_text(),
                  rating="", texts=_texts(body))
    assert "<table>" in html and "<td>100万t</td>" in html
    # 竖线不裸漏在段落里（表格单元格外）
    paras = [ln for ln in html.splitlines() if ln.strip().startswith("<p>")
             and "|" in ln]
    assert not paras, paras


def test_docx_renders_table(monkeypatch=None):
    """docx 渲染含表格块的正文不炸，且表格数正确（2 个：正文表 + 溯源附录）。"""
    from docx import Document
    from render import docx_report
    body = f"段一开头。\n{MD_TABLE}\n段二收尾。"
    out = Path("artifacts/_test_render_blocks/final.docx")
    docx_report.render_docx(_doc_one_text(body), {"title": "T"}, [], {}, {},
                            _spec_one_text(), rating="", kline_png=None,
                            out_path=out, texts=_texts(body))
    d = Document(str(out))
    assert len(d.tables) == 2       # 正文内嵌表 + 数据溯源附录
    cells = [c.text for c in d.tables[0].rows[1].cells]
    assert "储量" in cells and "100万t" in cells


def test_plain_text_unaffected():
    """无表格的正文字数/段落不受分块影响。"""
    from render.html_report import render
    body = "段一。\n段二。\n段三。"
    html = render(_doc_one_text(body), {"title": "T"}, [], {}, {},
                  _spec_one_text(), rating="", texts=_texts(body))
    assert html.count("<p>段") == 3


if __name__ == "__main__":
    test_split_blocks_table_and_text()
    test_html_renders_real_table_no_bare_pipes()
    test_docx_renders_table()
    test_plain_text_unaffected()
    print("test_render_blocks 全部通过")
