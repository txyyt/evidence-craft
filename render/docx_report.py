"""⑦ docx 渲染（规范简洁版）：与 html_report 同输入，黑白正式版式。

python-docx 代码生成（无模板文件依赖）。部门红头/图框等强版式走
"docx 模板文件 + 占位替换"路线，M9 按试点部门真实模板做。
题注规范：表题在上（表1 …），图题在下（图1 …），编号由渲染层统一排。
"""

from datetime import datetime
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from pipeline.sections import render_table_sec
from render.common import md_table_rows, split_md_blocks
from template_factory.schema import SpecV2

FONT = "微软雅黑"
_CN_NUM = "一二三四五六七八九十"


def _run(para, text: str, *, size: float = 10.5, bold: bool = False,
         color: str | None = None, center: bool = False) -> None:
    r = para.add_run(text)
    r.font.name = FONT
    r._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    r.font.size = Pt(size)
    r.font.bold = bold
    if color:
        r.font.color.rgb = RGBColor.from_string(color)
    if center:
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER


def _heading(doc: Document, text: str, size: float = 14) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_with_next = True   # 标题不孤悬页底
    _run(p, text, size=size, bold=True)


def _table_rows_layout(t) -> None:
    """行禁止跨页拆分；首行设为重复表头（续页自动带表头）。"""
    for i, row in enumerate(t.rows):
        tr_pr = row._tr.get_or_add_trPr()
        tr_pr.append(OxmlElement("w:cantSplit"))
        if i == 0:
            th = OxmlElement("w:tblHeader")
            th.set(qn("w:val"), "true")
            tr_pr.append(th)


def _add_table(doc: Document, rows: list[list[str]], small: bool = False) -> None:
    """rows 含表头行；Table Grid 样式（黑色细边框，黑白规范版式）。"""
    t = doc.add_table(rows=len(rows), cols=len(rows[0]))
    t.style = "Table Grid"
    _table_rows_layout(t)
    for i, row in enumerate(rows):
        for j, cell in enumerate(row):
            c = t.cell(i, j)
            c.text = ""
            p = c.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            _run(p, cell, size=8.5 if small else 10, bold=(i == 0))
    # 表后隔断段：压缩到 6pt，避免把后续内容挤出页尾（相邻表格靠它防合并）
    sp = doc.add_paragraph()
    sp.paragraph_format.space_after = Pt(0)
    sp.paragraph_format.line_spacing = Pt(6)


def _cn_no(i: int) -> str:
    """章节序号：一~十，之后用阿拉伯数字（章节多的报告不炸编号）。"""
    return _CN_NUM[i - 1] if 1 <= i <= len(_CN_NUM) else str(i)


def render_docx(doc: dict[str, Any], outline: dict[str, Any],
                sections: list[dict[str, Any]], forecast: dict[str, Any],
                risks: dict[str, Any], spec: SpecV2, rating: str,
                kline_png: str | None, out_path: Path,
                texts: list[dict[str, Any]] | None = None,
                notes: dict[str, str] | None = None,
                charts: dict[str, list[dict[str, str]]] | None = None) -> Path:
    texts_by_id = {t.get("section_id"): t for t in (texts or [])}
    notes = notes or {}
    charts = charts or {}
    d = Document()
    normal = d.styles["Normal"]
    normal.font.name = FONT
    normal.font.size = Pt(10.5)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    for sec in d.sections:
        sec.top_margin = sec.bottom_margin = Cm(2.2)
        sec.left_margin = sec.right_margin = Cm(2.4)

    meta = doc["meta"]
    # 标题区（黑白：标题居中 + 元信息一行 + 评级）
    p = d.add_paragraph()
    _run(p, outline["title"], size=16, bold=True, center=True)
    p = d.add_paragraph()
    from render.html_report import _clean_genknow, _data_cutoff
    _run(p, " ".join(x for x in (meta.get("stock", ""), meta.get("name", ""),
                                  meta.get("industry", "")) if x)
         + f"　｜　数据截至 {_data_cutoff(doc)}（语料文献口径）"
         + f"　｜　报告日期 {datetime.now():%Y-%m-%d}"
         + (f"　｜　评级 {rating}" if rating else ""),
        size=9, color="888888", center=True)

    fig_no = 0
    if kline_png:
        fig_no += 1
        p = d.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.keep_with_next = True   # 图与图题不分开
        p.add_run().add_picture(kline_png, width=Cm(15.5))
        p = d.add_paragraph()
        _run(p, f"图{fig_no}　近期大盘走势（上证指数，供行情参照）",
             size=9, color="888888", center=True)

    def _add_charts(jobs: list[dict[str, str]] | None) -> None:
        """图件（正文后内嵌或图件集），图号全文连续。"""
        nonlocal fig_no
        for j in jobs or []:
            fig_no += 1
            p = d.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.keep_with_next = True  # 图与图题不分开
            try:
                p.add_run().add_picture(j["png"], width=Cm(15.5))
            except Exception:  # noqa: BLE001 —— 图片缺失不炸渲染
                continue
            p = d.add_paragraph()
            _run(p, f"图{fig_no}　{j['caption']}",
                 size=9, color="888888", center=True)

    table_no = 0
    view_no = 0
    for i, sec in enumerate(spec.sections, 1):
        # 内嵌位无章节标题；subheading 即本节标题时同理（避免双重标题）
        if sec.inline or (sec.subheading and not sec.heading):
            pass
        else:
            _heading(d, sec.heading or f"{_cn_no(i)}、{sec.title}")
        if sec.subheading:
            p = d.add_paragraph()
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.keep_with_next = True
            _run(p, sec.subheading, size=12, bold=True)
        if sec.kind == "views":
            for s in sections:
                view_no += 1
                p = d.add_paragraph()
                label = (f"{sec.view_numbering}.{view_no}　"
                         if sec.view_numbering else f"{view_no}、")
                _run(p, f"{label}{s['heading']}", size=11, bold=True)
                p.paragraph_format.space_before = Pt(8)
                p.paragraph_format.keep_with_next = True  # 小标题与正文同页
                p = d.add_paragraph()
                _run(p, _clean_genknow(s["body"]))
        elif sec.kind == "table":
            rows, _note = md_table_rows(render_table_sec(doc, sec, spec))
            if rows:
                table_no += 1
                p = d.add_paragraph()
                p.paragraph_format.keep_with_next = True  # 表题与表格同页
                _run(p, f"表{table_no}　{sec.title}", size=9, bold=True, center=True)
                _add_table(d, rows)
            note_body = notes.get(sec.id, forecast.get("body", ""))
            if note_body:
                d.add_paragraph()
                _run(d.add_paragraph(), _clean_genknow(note_body))
        elif sec.kind == "risk":
            _run(d.add_paragraph(), _clean_genknow(risks.get("body", "")))
        elif sec.kind == "text":
            t = texts_by_id.get(sec.id)
            if t:
                # A1：正文分块——手写 markdown 表格块走表格渲染，其余按段落
                for bkind, btext in split_md_blocks(t["body"]):
                    if bkind == "table":
                        rows, _note = md_table_rows(btext)
                        if rows:
                            table_no += 1
                            p = d.add_paragraph()
                            p.paragraph_format.keep_with_next = True
                            _run(p, f"表{table_no}", size=9, bold=True, center=True)
                            _add_table(d, rows)
                        continue
                    for para in btext.split("\n"):
                        if para.strip():
                            _run(d.add_paragraph(), _clean_genknow(para.strip()))
        # 图件：figures 章节为图件集，text/views/risk 章节正文后内嵌
        # （范文形态：图随文走），图号全文连续
        _add_charts(charts.get(sec.id))

    # 数据溯源附录
    _heading(d, "附：数据溯源")
    _add_table(d, [["事实编号", "指标", "值", "来源", "截至"]]
               + [[f["id"], f["name"], f"{f['value']}{f['unit']}",
                   f["source"], f["as_of"]] for f in doc["facts"]], small=True)

    # 尾注与生成时间同段，避免最后一行孤悬新页
    tail = f" 生成时间 {datetime.now():%Y-%m-%d %H:%M}。"
    p = d.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    if spec.disclaimer:
        _run(p, spec.disclaimer, size=8.5, color="999999")
    _run(p, tail, size=8.5, color="999999")

    out_path.parent.mkdir(parents=True, exist_ok=True)
    d.save(out_path)
    return out_path
