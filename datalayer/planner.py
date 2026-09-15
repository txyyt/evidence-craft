"""意图驱动数据规划器（Phase B）。

用户给一段话写作意图（+ 可选本地资料文件夹），规划器产出一次性的数据
采集计划：检索本地语料库什么、联网搜什么、沿用哪些表格绑定——执行后
照常走"抽取即对账"→大纲→分节→judge 流水线。

设计边界：
- 规划器只产生"查询"，不产生事实——事实仍由 adapters 抽取并对账，
  本模块不碰任何数字
- LLM 规划失败（余额不足/网络/输出不合规）→ 规则兜底：沿用模板静态
  绑定的查询原样执行——系统永远可跑；plan.json 标注 mode=fallback
- 指定本地文件夹 → 现场建语料库（目录内容指纹缓存，未变化直接复用），
  动态 rag 绑定指向新片段库，模板静态 rag 绑定停用（换库不混库）
- plan.json 随运行落 artifacts：规划方式、全部查询、执行留痕——可审计
"""

import hashlib
import json
import re
from datetime import datetime
from pathlib import Path
from typing import Any

from datalayer.settings import settings

_RAG_ADAPTER, _WEB_ADAPTER = "rag_client", "web_search"
_DB_ADAPTER = "sqlite_query"
MAX_RAG, MAX_WEB, MAX_DB = 8, 6, 3


def _db_schemas() -> dict[str, str]:
    """settings.databases 声明的 SQLite 库 → {db_ref: "表(列 类型, ...)"文本}。
    只读表名+列名（PRAGMA），不取任何业务数据——schema 仅供规划器写查询。"""
    import sqlite3
    out: dict[str, str] = {}
    for ref, cfg in (settings.databases or {}).items():
        path = cfg.get("path")
        if not path:
            continue
        try:
            con = sqlite3.connect(settings.resolve(path))
            try:
                tables = [r[0] for r in con.execute(
                    "SELECT name FROM sqlite_master WHERE type='table'"
                    " AND name NOT LIKE 'sqlite_%'")]
                parts = []
                for t in tables:
                    cols = ", ".join(
                        f"{c[1]} {c[2]}"
                        for c in con.execute(f'PRAGMA table_info("{t}")'))
                    parts.append(f"{t}({cols})" if cols else t)
                if parts:
                    out[ref] = "；".join(parts)
            finally:
                con.close()
        except Exception:  # noqa: BLE001 —— 单库读失败不阻塞规划
            continue
    return out


def ensure_corpus(folder: str) -> dict[str, Any] | None:
    """把用户资料文件夹现场建为检索片段库（内容指纹缓存）。

    指纹 = 目录内 PDF 的（相对路径+大小+mtime）摘要；一致即复用已建库。
    返回 {dir, fragments_rel, n_fragments, n_files, fingerprint, rebuilt}；
    文件夹没有 PDF 时返回 None（F11：文件夹可能只带 Excel——不作为错误）。
    """
    src = Path(folder)
    if not src.is_dir():
        raise ValueError(f"资料文件夹不存在或不是目录：{folder}")
    entries = sorted(
        (str(p.relative_to(src)), p.stat().st_size, int(p.stat().st_mtime))
        for p in src.rglob("*.pdf") if p.stat().st_size > 0)
    if not entries:
        return None
    fp = hashlib.sha256(
        json.dumps(entries, ensure_ascii=False).encode("utf-8")).hexdigest()[:12]
    cache_dir = settings.resolve(f"data/corpus/_intent/{fp}")
    out = cache_dir / "fragments.json"
    # mock_fragments 由 rag adapter 按 settings 根解析，存项目相对 posix 路径
    rel = f"data/corpus/_intent/{fp}/fragments.json"
    if out.exists():
        index = json.loads((cache_dir / "index.json").read_text(encoding="utf-8"))
        return {"dir": str(cache_dir), "fragments_rel": rel,
                "n_fragments": index.get("n_fragments", 0),
                "n_files": index.get("n_files_ok", 0),
                "fingerprint": fp, "rebuilt": False}
    from datalayer.corpus import build
    index = build(src, out)
    return {"dir": str(cache_dir), "fragments_rel": rel,
            "n_fragments": index.get("n_fragments", 0),
            "n_files": index.get("n_files_ok", 0),
            "fingerprint": fp, "rebuilt": True}


# ---------- F11：用户资料文件夹的 Excel → 确定性 xlsx_table 绑定 ----------

def _num_ok(v: Any) -> bool:
    try:
        float(str(v).replace(",", ""))
        return True
    except (TypeError, ValueError):
        return False


def _unit_of(col_name: str) -> str:
    """unit 取列名中的括号内容（如「品位(%)」→ %），无括号为空。"""
    m = re.search(r"[（(]([^）)]{1,10})[）)]\s*$", col_name or "")
    return m.group(1) if m else ""


def xlsx_bindings(folder: str) -> tuple[list[dict[str, Any]],
                                        list[dict[str, Any]]]:
    """资料文件夹中每个 xlsx → 一条确定性 xlsx_table 绑定（F11）。
    规则：id/table_id=文件名词干、sheet=首个、id_column=首个非空列、
    value_columns 采样前 50 行识别数值列（unit 取列名括号）、
    table_columns 取全部列。
    返回 (bindings, metas)：bindings 直入 plan["file_bindings"]；
    metas=[{need, file, table_id, headers}] 供规划器提示词与回执覆盖判定。"""
    import openpyxl
    src = Path(folder)
    if not src.is_dir():
        raise ValueError(f"资料文件夹不存在或不是目录：{folder}")
    bindings: list[dict[str, Any]] = []
    metas: list[dict[str, Any]] = []
    for p in sorted(src.rglob("*.xlsx")):
        if p.name.startswith("~$"):
            continue
        try:
            wb = openpyxl.load_workbook(p, data_only=True, read_only=True)
            ws = wb[wb.sheetnames[0]]
            rows = list(ws.iter_rows(min_row=1, max_row=51, values_only=True))
            wb.close()
        except Exception as e:  # noqa: BLE001 —— 单个文件损坏不阻塞
            metas.append({"need": f"xlsx_{p.stem}"[:40], "file": p.name,
                          "table_id": "", "headers": [],
                          "error": f"读取失败：{e}"})
            continue
        if not rows:
            continue
        header = [str(c).strip() if c is not None else "" for c in rows[0]]
        if not any(header):
            continue
        id_idx = next((i for i, h in enumerate(header) if h), 0)
        id_column = header[id_idx] or f"列{id_idx + 1}"
        sample = [r for r in rows[1:51]
                  if any(c is not None and str(c).strip() != "" for c in r)]

        def numeric(col_i: int) -> bool:
            vals = [r[col_i] for r in sample
                    if col_i < len(r) and r[col_i] not in (None, "", "-")]
            if not vals:
                return False
            return sum(1 for v in vals if _num_ok(v)) / len(vals) >= 0.6

        value_columns = []
        for i, h in enumerate(header):
            if not h or i == id_idx or not numeric(i):
                continue
            key = re.sub(r"[^\w]+", "_", h)[:30].strip("_") or f"c{i}"
            value_columns.append({"column": h, "key": key,
                                  "unit": _unit_of(h)})
        stem = re.sub(r"[^\w\u4e00-\u9fff-]+", "_", p.stem)[:40].strip("_") \
            or f"xlsx_{len(bindings) + 1}"
        need = f"xlsx_{stem}"
        bindings.append({
            "need": need, "adapter": "xlsx_table",
            "params": {"path": str(p), "header_row": 1,
                       "id_prefix": stem, "id_column": id_column,
                       "name_template": "{" + id_column + "}",
                       "value_columns": value_columns,
                       "table_columns": [h for h in header if h],
                       "table_id": stem}})
        metas.append({"need": need, "file": p.name, "table_id": stem,
                      "headers": [h for h in header if h],
                      "n_value_columns": len(value_columns)})
    return bindings, metas


def _structure_lines(spec: Any) -> str:
    """报告结构概览：章节 + 观点槽位职责（规划器据此判断需要什么数据）。"""
    lines = []
    for s in spec.sections:
        line = f"- {s.id}（{s.title}，{s.kind}）"
        if s.view_slots:
            line += "：" + "；".join(
                f"{v.id}={v.brief[:40]}" for v in s.view_slots)
        elif s.style:
            line += "：" + s.style[:40].replace("\n", " ")
        lines.append(line)
    return "\n".join(lines)


def _static_bindings(sources: dict[str, Any]) -> tuple[list[dict], list[dict],
                                                       list[dict]]:
    rag, web, other = [], [], []
    for b in sources.get("bindings") or []:
        key = b.get("adapter")
        if key == _RAG_ADAPTER:
            rag.append(b)
        elif key == _WEB_ADAPTER:
            web.append(b)
        else:
            other.append(b)
    return rag, web, other


_PLAN_SYSTEM = """你是矿产资源报告的数据规划师。根据写作意图与报告结构，为本次
报告生成数据采集计划：确定检索本地语料库的关键词、联网搜索的查询词，以及
（数据库与意图相关时）对结构化数据库的查询。只规划"查询"，不编造任何数据。
查询要具体（含矿种/主题词/年份），中文为主。意图涉及多个主体（如多个矿种、
多个地区）时按主体拆分查询。数据库查询只写单条 SELECT（可用 :param 命名参数，
参数值放 query_params，值可写 "$project" 这类运行参数引用），不带 LIMIT（系统
自动限量）。只输出 JSON。"""

_PLAN_USER_TMPL = """【写作意图】
{intent}

【报告结构】（各章节需要数据支撑）
{structure}

【数据需求清单】（树上各节声明的语义需求；逐条判定 covered/search/gap：
covered=已有语料检索可支撑，search=需联网搜索，gap=计划无来源）
{needs}

【可用数据库】（仅当库中表与意图相关时规划 db 查询；value_columns 的 key 用英文）
{databases}

【可用运行参数】（query_params 的值可引用，写成 "$参数名" 字符串）
{run_params}

【模板现有查询】（可沿用、改写或增删；need 命名保持 corpus_*/web_*/db_* 风格）
{static}

任务：输出 JSON：
{{"subject": "报告主题短语", "focus": "本次报告内容要点的一句话归纳（供写作选材聚焦）",
"rag": [{{"need": "corpus_x", "query": "检索关键词", "top_k": 5}}],
"web": [{{"need": "web_x", "query": "联网搜索查询词", "top_k": 6, "fetch_pages": 3}}],
"db": [{{"need": "db_x", "db_ref": "库名", "query": "SELECT 列 FROM 表 WHERE ...",
        "query_params": {{"project": "$project"}}, "id_column": "主键列",
        "name_template": "{{列名}} 描述",
        "value_columns": [{{"column": "数值列", "key": "英文键", "unit": "单位"}}],
        "as_of_column": "日期列（可选）", "table_columns": ["展示列..."],
        "table_id": "表格id（可选，配 table_columns 时生成整表）"}}],
"tables_kept": ["需沿用的表格/数据库绑定 need 名"],
"needs_coverage": [{{"need": "需求清单原句", "status": "covered|search|gap"}}]}}
rag 不超过 {max_rag} 条、web 不超过 {max_web} 条、db 不超过 {max_db} 条。
【可用数据库】为空时 db 输出空数组。needs_coverage 必须覆盖需求清单的每一句。"""


def _chart_table_ids(spec: Any) -> set[str]:
    """spec 中被图件模板引用的表格 id（table:<tid>）——图表数据源必须保留，
    规划器不得按"意图相关性"剔除，否则渲染期缺表少图。"""
    out: set[str] = set()
    for s in spec.sections:
        for c in s.charts:
            if c.source.startswith("table:"):
                out.add(c.source.split(":", 1)[1])
    return out


def _clamp_plan(out: dict[str, Any], other: list[dict],
                keep_ids: set[str] | None = None,
                db_schemas: dict[str, str] | None = None,
                needs: list[str] | None = None
                ) -> tuple[dict[str, Any], list[str]]:
    """LLM 计划清洗：条数/类型/取值范围收敛，非法项丢弃；
    keep_ids 中的表格绑定强制沿用（图表数据源，不受意图取舍影响）；
    db 计划逐条过 sanitize_sql 安全约束，违规拒绝并记 warning。
    返回 (计划 dict, 拒绝/告警列表)。"""
    def _s(v: Any, hi: int = 120) -> str:
        return str(v or "").strip()[:hi]

    rag = []
    for q in (out.get("rag") or [])[:MAX_RAG]:
        query = _s(q.get("query") if isinstance(q, dict) else "")
        if query:
            rag.append({"need": _s(q.get("need"), 40) or f"corpus_{len(rag)}",
                        "query": query,
                        "top_k": min(max(int(q.get("top_k") or 5), 2), 10)})
    web = []
    for q in (out.get("web") or [])[:MAX_WEB]:
        query = _s(q.get("query") if isinstance(q, dict) else "")
        if query:
            web.append({"need": _s(q.get("need"), 40) or f"web_{len(web)}",
                        "query": query,
                        "top_k": min(max(int(q.get("top_k") or 6), 2), 10),
                        "fetch_pages": min(max(int(q.get("fetch_pages") or 3), 1), 5)})
    # db 计划：安全约束代码级强制（单条 SELECT / 禁多语句 / 自动 LIMIT）
    from datalayer.adapters.database import sanitize_sql
    dbs: list[dict[str, Any]] = []
    warnings: list[str] = []
    for q in (out.get("db") or [])[:MAX_DB]:
        if not isinstance(q, dict):
            continue
        db_ref = _s(q.get("db_ref"), 40)
        query = _s(q.get("query"), 2000)
        if db_ref not in (db_schemas or {}):
            warnings.append(f"db 规划被拒：未知数据库引用 {db_ref!r}")
            continue
        id_column = _s(q.get("id_column"), 80)
        if not id_column:
            warnings.append(f"db 规划被拒（{db_ref}）：缺 id_column")
            continue
        try:
            clean_query = sanitize_sql(query)
        except ValueError as e:
            warnings.append(f"db 查询被拒绝（{db_ref}）：{e} —— {query[:80]}")
            continue
        item: dict[str, Any] = {
            "need": _s(q.get("need"), 40) or f"db_{len(dbs)}",
            "db_ref": db_ref, "query": clean_query,
            "id_column": id_column,
            "name_template": _s(q.get("name_template"), 120)
            or f"{{{id_column}}}",
            "value_columns": []}
        qparams = q.get("query_params")
        if isinstance(qparams, dict) and qparams:
            # 参数值一律字符串化（执行侧绑定参数，不拼接进 SQL 文本）
            item["query_params"] = {str(k)[:40]: str(v)[:200]
                                    for k, v in list(qparams.items())[:8]}
        for vc in (q.get("value_columns") or [])[:6]:
            if isinstance(vc, dict) and vc.get("column"):
                item["value_columns"].append(
                    {"column": _s(vc.get("column"), 80),
                     "key": _s(vc.get("key"), 40) or "value",
                     "unit": _s(vc.get("unit"), 20)})
        if q.get("as_of_column"):
            item["as_of_column"] = _s(q.get("as_of_column"), 80)
        cols = [_s(c, 80) for c in (q.get("table_columns") or [])[:8] if c]
        if cols:
            item["table_columns"] = cols
            item["table_id"] = _s(q.get("table_id"), 40) or item["need"]
        dbs.append(item)
    need_ids = {b.get("need") for b in other}
    kept = [t for t in (out.get("tables_kept") or [])
            if isinstance(t, str) and t in need_ids]
    for b in other:                # 图表数据源强制保留
        tid = b.get("table_id") or (b.get("params") or {}).get("table_id")
        if keep_ids and tid in keep_ids and b.get("need") not in kept:
            kept.append(b["need"])
    for b in other:                # 静态 db 绑定不可剔除
        if b.get("adapter") == _DB_ADAPTER and b.get("need") not in kept:
            kept.append(b["need"])
    if not kept:
        kept = [b.get("need") for b in other]
    return {"subject": _s(out.get("subject"), 40) or "报告主题",
            "focus": _s(out.get("focus"), 200),
            "rag": rag, "web": web, "db": dbs, "tables_kept": kept,
            "needs_coverage": _clamp_coverage(out.get("needs_coverage"), needs)}, \
        warnings


def _clamp_coverage(raw: Any, needs: list[str] | None) -> list[dict[str, str]]:
    """needs_coverage 清洗：status 白名单；needs 给出时保证逐条覆盖（缺失→gap）。"""
    if not needs:
        return []
    by_need = {}
    for it in raw or []:
        if isinstance(it, dict) and it.get("need"):
            status = it.get("status")
            by_need[str(it["need"])[:80]] = status if status in (
                "covered", "search", "gap") else "gap"
    out = []
    for n in needs:
        status = by_need.get(n) or by_need.get(n[:80])
        out.append({"need": n, "status": status or "gap"})
    return out


def _fallback_plan(intent: str, rag_static: list[dict], web_static: list[dict],
                   other: list[dict], error: str) -> dict[str, Any]:
    """规则兜底：静态绑定查询原样沿用（保真不保聚焦），plan 标注 fallback。
    兜底不臆造 db 查询——静态 db 绑定经 tables_kept 原样执行。"""
    return {
        "mode": "fallback", "error": error,
        "subject": intent[:20] or "报告主题",
        "focus": intent,
        "rag": [{"need": b.get("need"), "query": (b.get("params") or {}).get("query", ""),
                 "top_k": (b.get("params") or {}).get("top_k", 5)}
                for b in rag_static],
        "web": [{"need": b.get("need"),
                 "query": (b.get("params") or {}).get("query", ""),
                 "top_k": (b.get("params") or {}).get("top_k", 6),
                 "fetch_pages": (b.get("params") or {}).get("fetch_pages", 3)}
                for b in web_static],
        "db": [],
        "tables_kept": [b.get("need") for b in other],
    }


def make_plan(spec: Any, sources: dict[str, Any], intent: str,
              folder: str | None = None,
              needs: list[str] | None = None) -> dict[str, Any]:
    """意图 + 结构模板 + 静态绑定 + 数据需求清单 → 采集计划。
    needs（树场景）：各节 data_needs 汇总；给出时计划附 needs_coverage 三态回执。"""
    intent = (intent or "").strip()
    if not intent and not folder and not needs:
        raise ValueError("意图（--intent）与资料文件夹（--folder）与数据需求"
                         "（needs）至少给一个")
    rag_static, web_static, other = _static_bindings(sources)
    chart_kept = _chart_table_ids(spec)

    corpus_info = None
    file_bindings: list[dict[str, Any]] = []
    file_metas: list[dict[str, Any]] = []
    if folder:
        corpus_info = ensure_corpus(folder)
        file_bindings, file_metas = xlsx_bindings(folder)

    plan: dict[str, Any] = {
        "created_at": datetime.now().isoformat(timespec="seconds"),
        "intent": intent, "corpus": corpus_info, "mode": "llm",
    }
    if file_bindings:
        plan["file_bindings"] = file_bindings
        plan["files"] = file_metas
    try:
        from pipeline.llm import chat_json, tier_for
        schemas = _db_schemas()
        static_lines = "\n".join(
            f"- rag need={b.get('need')} query={(b.get('params') or {}).get('query', '')}"
            for b in rag_static)
        static_lines += "\n" + "\n".join(
            f"- web need={b.get('need')} query={(b.get('params') or {}).get('query', '')}"
            for b in web_static)
        static_lines += "\n" + "\n".join(
            f"- table/db need={b.get('need')} adapter={b.get('adapter')}"
            for b in other)
        # F11：告知规划器用户提供的表格（选材时可直接判定 covered）
        if file_metas:
            static_lines += "\n" + "\n".join(
                f"- 用户表格 need={m['need']} 文件={m['file']} "
                f"列={('、'.join(m['headers']))[:80]}"
                for m in file_metas if not m.get("error"))
        db_lines = "\n".join(f"- {ref}: {txt}" for ref, txt in schemas.items())
        run_params = "、".join(f"${k}（{v}）"
                               for k, v in (sources.get("params_schema") or {}).items())
        out = chat_json(
            _PLAN_SYSTEM,
            _PLAN_USER_TMPL.format(
                intent=intent or "（无文字意图，以资料文件夹为主题生成综述报告）",
                structure=_structure_lines(spec),
                needs="\n".join(f"- {n}" for n in (needs or [])) or "（无）",
                static=static_lines.strip(),
                databases=db_lines.strip() or "（无已配置数据库）",
                run_params=run_params.strip() or "（无）",
                max_rag=MAX_RAG, max_web=MAX_WEB, max_db=MAX_DB),
            schema_hint="只输出一个合法 JSON 对象："
                        "subject/focus/rag/web/db/tables_kept/needs_coverage。",
            tier=tier_for("extract"))
        cleaned, warnings = _clamp_plan(out, other, chart_kept, schemas,
                                        needs=needs)
        if not cleaned["rag"] and not cleaned["web"] and not cleaned["db"]:
            raise ValueError("规划器未产出任何查询")
        plan.update(cleaned)
        if warnings:
            plan["warnings"] = warnings
    except Exception as e:  # noqa: BLE001 —— 规划失败兜底，不阻塞报告生成
        plan.update(_fallback_plan(intent, rag_static, web_static, other,
                                   f"{type(e).__name__}: {e}"))
        if needs:                       # 兜底：全部标 gap 交用户裁决
            plan["needs_coverage"] = [{"need": n, "status": "gap"} for n in needs]
    # F11：确定性回执升级——需求词元命中用户表格文件名/列名 → 计已覆盖
    if file_metas:
        hay = " ".join((m.get("file") or "") + " " + " ".join(m.get("headers") or [])
                       for m in file_metas)
        for c in plan.get("needs_coverage") or []:
            if c.get("status") != "gap":
                continue
            toks = [t for t in str(c.get("need", "")).split() if len(t) >= 2]
            if toks and any(t in hay for t in toks):
                c["status"] = "covered"
                c["source"] = "xlsx"
    # 换库：意图文件夹建了新语料库 → 全部 rag 查询指向新库（不与模板库混用）
    if corpus_info:
        plan["corpus_switched"] = True
    return plan


def plan_to_bindings(plan: dict[str, Any], sources: dict[str, Any]) -> list[dict]:
    """计划 → 动态绑定列表（sources.yaml 同 schema，registry 直接执行）。"""
    rag_static = {b.get("need"): b for b in sources.get("bindings") or []
                  if b.get("adapter") == _RAG_ADAPTER}
    web_static = {b.get("need"): b for b in sources.get("bindings") or []
                  if b.get("adapter") == _WEB_ADAPTER}
    corpus_rel = (plan.get("corpus") or {}).get("fragments_rel")
    bindings: list[dict] = []
    for q in plan.get("rag") or []:
        params: dict[str, Any] = {"query": q["query"], "top_k": q.get("top_k", 5)}
        if corpus_rel:
            params["mock_fragments"] = corpus_rel
        else:                       # 未换库：沿用静态绑定的片段库指向
            base = (rag_static.get(q.get("need")) or
                    next(iter(rag_static.values()), None))
            if base:
                base_params = base.get("params") or {}
                if base_params.get("mock_fragments"):
                    params["mock_fragments"] = base_params["mock_fragments"]
        if params.get("mock_fragments") or corpus_rel is None:
            bindings.append({"need": q["need"], "adapter": _RAG_ADAPTER,
                             "params": params})
    for q in plan.get("web") or []:
        params = {"query": q["query"], "top_k": q.get("top_k", 6),
                  "fetch_pages": q.get("fetch_pages", 3), "cache_ttl_h": 12}
        bindings.append({"need": q["need"], "adapter": _WEB_ADAPTER,
                         "params": params})
    # 数据库动态绑定：计划里的查询全文落 plan.json（可审计），执行侧再过
    # 一道 sanitize_sql（adapter 内代码强制，双保险）
    for q in plan.get("db") or []:
        params: dict[str, Any] = {"db_ref": q["db_ref"], "query": q["query"],
                                  "id_prefix": q.get("need") or "db",
                                  "id_column": q["id_column"],
                                  "name_template": q.get("name_template")
                                  or f"{{{q['id_column']}}}",
                                  "value_columns": q.get("value_columns") or []}
        if q.get("query_params"):
            params["query_params"] = q["query_params"]
        if q.get("as_of_column"):
            params["as_of_column"] = q["as_of_column"]
        if q.get("table_columns"):
            params["table_columns"] = q["table_columns"]
            params["table_id"] = q.get("table_id") or q["need"]
        bindings.append({"need": q["need"], "adapter": _DB_ADAPTER,
                         "params": params})
    # 表格等静态绑定按计划沿用（planner 可按意图剔除无关表）
    kept = set(plan.get("tables_kept") or [])
    for b in sources.get("bindings") or []:
        if b.get("adapter") not in (_RAG_ADAPTER, _WEB_ADAPTER) \
                and b.get("need") in kept:
            bindings.append(b)
    # F11：用户资料文件夹的 xlsx 确定性绑定（事实层+表格层）
    bindings.extend(plan.get("file_bindings") or [])
    return bindings
