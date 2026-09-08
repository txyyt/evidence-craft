"""local_file 类 adapter：xlsx 声明式映射 + inbox 人工投喂目录。

- 来源指纹：文件名 + 内容 md5 前 8 位写入 fact.source，改文件即换指纹
- path 可以是目录：配合 pattern 取 mtime 最新的文件（inbox 投喂习惯：
  同名新版本直接覆盖/累加，系统永远吃最新一份）
"""

import hashlib
from pathlib import Path
from typing import Any

from datalayer.adapters.base import (AdapterResult, SourceAdapter,
                                     rows_to_facts, rows_to_table)
from datalayer.settings import settings


def _resolve_path(path: str, pattern: str | None) -> Path:
    p = settings.resolve(path)
    if p.is_dir():
        if not pattern:
            raise ValueError(f"inbox 目录 {p} 需要 pattern 参数")
        files = sorted(p.glob(pattern), key=lambda f: f.stat().st_mtime)
        if not files:
            raise ValueError(f"inbox 目录 {p} 无匹配 {pattern} 的文件")
        return files[-1]
    return p


def _read_rows(path: Path, sheet: str | None, header_row: int) -> list[dict[str, Any]]:
    import openpyxl
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    ws = wb[sheet] if sheet else wb[wb.sheetnames[0]]
    rows = list(ws.iter_rows(min_row=header_row, values_only=True))
    if not rows:
        return []
    header = [str(c).strip() if c is not None else "" for c in rows[0]]
    out = []
    for raw in rows[1:]:
        if all(c is None or str(c).strip() == "" for c in raw):
            continue
        out.append({header[i]: raw[i] for i in range(len(header)) if header[i]})
    return out


def _fingerprint(path: Path) -> str:
    digest = hashlib.md5(path.read_bytes()).hexdigest()[:8]
    return f"local_file:{path.name}#{digest}"


class XlsxAdapter(SourceAdapter):
    """xlsx 台账 → 事实。params：
    path/sheet/header_row：文件定位；path 为目录时按 pattern 取最新
    id_prefix/id_column/name_template/value_columns/as_of_column：声明式映射
    table_columns：可选，同时产出 collections["tables"][table_id] 供
    generic_rows 渲染器出表格
    """
    key = "xlsx_table"
    kind = "local_file"

    def fetch(self, params: dict[str, Any]) -> AdapterResult:
        path = _resolve_path(params["path"], params.get("pattern"))
        rows = _read_rows(path, params.get("sheet"),
                          int(params.get("header_row", 1)))
        if not rows:
            raise ValueError(f"{path} 未读到数据行")
        source = _fingerprint(path)
        facts = rows_to_facts(
            rows, source=source,
            id_prefix=params["id_prefix"], id_column=params["id_column"],
            name_template=params["name_template"],
            value_columns=params.get("value_columns") or [],
            as_of_column=params.get("as_of_column"))
        result = AdapterResult(facts=facts)
        if params.get("table_columns"):
            tid = params.get("table_id") or params["id_prefix"]
            result.collections["tables"] = {
                tid: rows_to_table(rows, params["table_columns"], tid)}
        return result
